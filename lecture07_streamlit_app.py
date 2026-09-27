
# INKOR FMCG Lecture 07 - Fulfilment Data Preparation Tool
# Combined Streamlit application.
# Core classification logic and StoreRobo mapping are preserved from the
# two validated source scripts supplied for the course.

import tempfile
import csv
import streamlit as st
import sys
import re
from pathlib import Path
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

COURIER = "DTDC India"

PAYMENT_HOURS = 4
TRACKING_HOURS = 8
DELIVERY_HOURS = 16
REFUND_HOURS = 24

REQUIRED_COLUMNS = {
    "Name",
    "Financial Status",
    "Fulfillment Status",
    "Currency",
    "Created at",
    "Lineitem quantity",
    "Lineitem name",
    "Lineitem sku",
    "Lineitem requires shipping",
    "Notes",
    "Cancelled at",
    "Refunded Amount",
    "Outstanding Balance",
    "Total",
    "Location",
    "Source",
}


# ============================================================
# HELPERS
# ============================================================

def txt(v):
    if pd.isna(v):
        return ""
    return str(v).strip()


def norm(v):
    return txt(v).lower().replace(" ", "_")


def num(v):
    x = pd.to_numeric(v, errors="coerce")
    return 0.0 if pd.isna(x) else float(x)


def dt(v):
    x = pd.to_datetime(v, errors="coerce")
    return None if pd.isna(x) else x.to_pydatetime()


def true_value(v):
    return norm(v) in {"true", "yes", "1", "y"}


def clean_order(v):
    return re.sub(r"[^A-Za-z0-9]", "", txt(v))


def display_dt(d):
    return d.strftime("%Y-%m-%d %H:%M:%S") if d else ""


def timestamp_id(d):
    return d.strftime("%Y%m%d%H%M") if d else "UNKNOWN"


def unique_nonblank(series):
    result = []

    for v in series:
        s = txt(v)
        if s and s not in result:
            result.append(s)

    return result


def first_nonblank(series):
    values = unique_nonblank(series)
    return values[0] if values else ""


def read_csv_file(path):
    return pd.read_csv(
        path,
        dtype=str,
        keep_default_na=False
    )


# ============================================================
# VALIDATE FILE
# ============================================================

def validate_columns(df, description):

    missing = sorted(
        REQUIRED_COLUMNS - set(df.columns)
    )

    if missing:
        raise ValueError(
            f"{description} is missing columns: "
            + ", ".join(missing)
        )


# ============================================================
# GET UNIQUE ORDER NAMES
# ============================================================

def order_names(df):

    return {
        txt(x)
        for x in df["Name"]
        if txt(x)
    }


# ============================================================
# ORDER RECORD
# ============================================================

def make_order_record(name, g, cancelled_orders, return_orders):

    created = dt(
        first_nonblank(g["Created at"])
    )

    fin = norm(
        first_nonblank(g["Financial Status"])
    )

    ful_raw = first_nonblank(
        g["Fulfillment Status"]
    )

    ful = norm(ful_raw) if ful_raw else "unfulfilled"

    cancelled_at = first_nonblank(
        g["Cancelled at"]
    )

    # External cancelled-order file is authoritative
    in_cancel_file = name in cancelled_orders

    cancelled = (
        bool(cancelled_at)
        or in_cancel_file
    )

    # External return-request file is authoritative
    return_requested = name in return_orders

    refund_amount = max(
        [num(v) for v in g["Refunded Amount"]]
        or [0.0]
    )

    outstanding = num(
        first_nonblank(g["Outstanding Balance"])
    )

    total = num(
        first_nonblank(g["Total"])
    )

    currency = first_nonblank(
        g["Currency"]
    )

    notes = first_nonblank(
        g["Notes"]
    )

    location = first_nonblank(
        g["Location"]
    )

    source = first_nonblank(
        g["Source"]
    )

    shipping_values = [
        true_value(v)
        for v in g["Lineitem requires shipping"]
        if txt(v)
    ]

    shipping_required = any(
        shipping_values
    )

    mixed_shipping = (
        bool(shipping_values)
        and any(shipping_values)
        and not all(shipping_values)
    )

    incorrect_order = total <= 0

    payment_pending = (
        fin in {"pending", "partially_paid"}
        and outstanding > 0
    )

    # COURSE DATASET RULE
    refund_pending = outstanding < 0

    refund_completed = (
        refund_amount > 0
    )

    fully_refunded = (
        fin == "refunded"
        and refund_amount > 0
        and outstanding == 0
    )

    # Cancelled + refunded + no outstanding work
    terminal_cancelled = (
        cancelled
        and fully_refunded
    )

    needs_fulfilment = (
        ful != "fulfilled"
        and not cancelled
    )

    return {
        "order": name,
        "g": g,

        "created": created,
        "fin": fin,
        "ful": ful,

        "cancelled": cancelled,
        "cancelled_at": cancelled_at,
        "in_cancel_file": in_cancel_file,

        "return_requested": return_requested,

        "refund_amount": refund_amount,
        "outstanding": outstanding,
        "total": total,

        "currency": currency,
        "notes": notes,
        "location": location,
        "source": source,

        "shipping_required": shipping_required,
        "mixed_shipping": mixed_shipping,

        "incorrect_order": incorrect_order,

        "payment_pending": payment_pending,
        "refund_pending": refund_pending,
        "refund_completed": refund_completed,
        "fully_refunded": fully_refunded,

        "terminal_cancelled": terminal_cancelled,

        "needs_fulfilment": needs_fulfilment,
    }


# ============================================================
# EXCEPTION
# ============================================================

def add_exception(rows, r, code, reason, action):

    rows.append({
        "Order": r["order"],
        "Financial Status": r["fin"],
        "Fulfillment Status": r["ful"],
        "Outstanding Balance": r["outstanding"],
        "Refunded Amount": r["refund_amount"],
        "Total": r["total"],
        "Cancelled": "Yes" if r["cancelled"] else "No",
        "Return Requested":
            "Yes" if r["return_requested"] else "No",
        "Notes": r["notes"],
        "Exception Code": code,
        "Exception Type": reason,
        "Required Action": action,
    })


# ============================================================
# BUILD WORKBOOK
# ============================================================

def build_workbook(
    main_csv,
    return_csv,
    cancelled_csv,
    output_path=None
):

    # --------------------------------------------------------
    # READ THREE FILES
    # --------------------------------------------------------

    main_df = read_csv_file(main_csv)
    return_df = read_csv_file(return_csv)
    cancelled_df = read_csv_file(cancelled_csv)

    validate_columns(
        main_df,
        "Main order export"
    )

    validate_columns(
        return_df,
        "Return-request export"
    )

    validate_columns(
        cancelled_df,
        "Cancelled-order export"
    )

    return_orders = order_names(
        return_df
    )

    cancelled_orders = order_names(
        cancelled_df
    )

    main_orders = order_names(
        main_df
    )

    # --------------------------------------------------------
    # CHECK REFERENCE FILES
    # --------------------------------------------------------

    unknown_returns = (
        return_orders - main_orders
    )

    unknown_cancelled = (
        cancelled_orders - main_orders
    )

    if unknown_returns:
        print(
            "WARNING: Return-request file contains "
            "orders absent from main export:"
        )
        print(
            ", ".join(sorted(unknown_returns))
        )

    if unknown_cancelled:
        print(
            "WARNING: Cancelled-order file contains "
            "orders absent from main export:"
        )
        print(
            ", ".join(sorted(unknown_cancelled))
        )

    # --------------------------------------------------------
    # OUTPUT SHEETS
    # --------------------------------------------------------

    sheets = {

        "00_Incorrect_Orders": [],

        "01_No_Action_Required": [],

        "02_Cancelled_Orders": [],

        "03_Return_Requested": [],

        "04_Payment_Received": [],

        "05_Refund_Processed": [],

        "06_Refund_Completed": [],

        "07_Fulfilment": [],

        "08_Tracking": [],

        "09_Delivery": [],

        "10_Import_Error_Resolution": [],

        "11_Exceptions": [],
    }

    # --------------------------------------------------------
    # UNIQUE ORDERS
    # --------------------------------------------------------

    records = []

    for name, g in main_df.groupby(
        "Name",
        sort=False,
        dropna=False
    ):

        name = txt(name)

        if not name:
            continue

        records.append(
            make_order_record(
                name,
                g,
                cancelled_orders,
                return_orders
            )
        )

    # ========================================================
    # PROCESS ORDERS
    # ========================================================

    for r in records:

        created = r["created"]

        payment_at = (
            created + timedelta(
                hours=PAYMENT_HOURS
            )
            if created else None
        )

        tracking_at = (
            created + timedelta(
                hours=TRACKING_HOURS
            )
            if created else None
        )

        delivery_at = (
            created + timedelta(
                hours=DELIVERY_HOURS
            )
            if created else None
        )

        refund_at = (
            created + timedelta(
                hours=REFUND_HOURS
            )
            if created else None
        )

        tracking_number = (
            f"DTDC"
            f"{clean_order(r['order'])}"
            f"{timestamp_id(tracking_at)}"
        )

        # ====================================================
        # 00 INCORRECT ORDER
        # ====================================================

        if r["incorrect_order"]:

            sheets["00_Incorrect_Orders"].append({
                "Order": r["order"],
                "Total": r["total"],
                "Currency": r["currency"],
                "Financial Status": r["fin"],
                "Fulfillment Status": r["ful"],
                "Outstanding Balance": r["outstanding"],
                "Refunded Amount": r["refund_amount"],
                "Incorrect Order": "Yes",
                "Cancel": "Yes",
                "Refund": "Yes",
                "Archive": "Yes",
                "Required Action":
                    "Cancel and archive incorrect order",
            })

            # Remove from every later operational sheet
            continue

        # ====================================================
        # 01 TERMINAL CANCELLED ORDERS
        # ====================================================
        #
        # Cancelled + fully refunded + outstanding = 0
        #
        # No payment / fulfilment / tracking / delivery needed.
        # ====================================================

        if r["terminal_cancelled"]:

            sheets["01_No_Action_Required"].append({
                "Order": r["order"],
                "Reason":
                    "Cancelled and refund completed",
                "Financial Status": r["fin"],
                "Fulfillment Status": r["ful"],
                "Refunded Amount": r["refund_amount"],
                "Outstanding Balance": r["outstanding"],
                "Required Processing": "None",
                "Archive": "Yes",
            })

            sheets["02_Cancelled_Orders"].append({
                "Order": r["order"],
                "Cancelled": "Yes",
                "Financial Status": r["fin"],
                "Refunded Amount": r["refund_amount"],
                "Outstanding Balance": r["outstanding"],
                "Status":
                    "Cancellation and refund completed",
                "Archive": "Yes",
            })

            # Completely eliminate from remaining workflow
            continue

        # ====================================================
        # 02 CANCELLED BUT NOT TERMINAL
        # ====================================================

        if r["cancelled"]:

            sheets["02_Cancelled_Orders"].append({
                "Order": r["order"],
                "Cancelled": "Yes",
                "Financial Status": r["fin"],
                "Fulfillment Status": r["ful"],
                "Refunded Amount": r["refund_amount"],
                "Outstanding Balance": r["outstanding"],
                "Status":
                    "Cancellation requires reconciliation",
                "Archive": "No",
            })

            # Cancelled orders must NOT go to normal
            # fulfilment/tracking/delivery.
            #
            # Refund/payment state is handled below.

        # ====================================================
        # 03 RETURN REQUESTED
        # ====================================================

        if r["return_requested"]:

            sheets["03_Return_Requested"].append({
                "Order": r["order"],
                "Financial Status": r["fin"],
                "Fulfillment Status": r["ful"],
                "Total": r["total"],
                "Refunded Amount": r["refund_amount"],
                "Outstanding Balance": r["outstanding"],
                "Notes": r["notes"],
                "Return Requested": "Yes",
            })

        # ====================================================
        # 04 PAYMENT RECEIVED
        # ====================================================
        #
        # Do not treat return-request or cancelled orders as
        # ordinary payment collection work.
        # ====================================================

        if (
            r["payment_pending"]
            and not r["refund_pending"]
            and not r["cancelled"]
            and not r["return_requested"]
        ):

            sheets["04_Payment_Received"].append({
                "Order": r["order"],
                "Previous Payment Status": r["fin"],
                "Pending Amount": r["outstanding"],
                "Payment Received Amount": r["outstanding"],
                "Currency": r["currency"],
                "Payment Received At":
                    display_dt(payment_at),
                "Payment Reference":
                    (
                        f"PAY-"
                        f"{clean_order(r['order'])}-"
                        f"{timestamp_id(payment_at)}"
                    ),
                "Resolved Financial Status": "paid",
                "Resolved Outstanding Balance": 0,
            })

        # ====================================================
        # 05 REFUND PROCESSED
        # ====================================================

        if r["refund_pending"]:

            refund_pending_amount = abs(
                r["outstanding"]
            )

            sheets["05_Refund_Processed"].append({
                "Order": r["order"],
                "Return Requested":
                    "Yes"
                    if r["return_requested"]
                    else "No",
                "Cancelled":
                    "Yes"
                    if r["cancelled"]
                    else "No",
                "Refund Pending Amount":
                    refund_pending_amount,
                "Refund Processed Amount":
                    refund_pending_amount,
                "Currency": r["currency"],
                "Refund Processed At":
                    display_dt(refund_at),
                "Refund Reference":
                    (
                        f"REF-"
                        f"{clean_order(r['order'])}-"
                        f"{timestamp_id(refund_at)}"
                    ),
                "Refund Note":
                    (
                        r["notes"]
                        if r["notes"]
                        else
                        "Historical refund reconciliation"
                    ),
                "Resolved Outstanding Balance": 0,
            })

        # ====================================================
        # 06 EXISTING COMPLETED REFUND
        # ====================================================

        if r["refund_completed"]:

            sheets["06_Refund_Completed"].append({
                "Order": r["order"],
                "Return Requested":
                    "Yes"
                    if r["return_requested"]
                    else "No",
                "Cancelled":
                    "Yes"
                    if r["cancelled"]
                    else "No",
                "Refunded Amount":
                    r["refund_amount"],
                "Currency":
                    r["currency"],
                "Financial Status":
                    r["fin"],
                "Outstanding Balance":
                    r["outstanding"],
                "Order Note":
                    r["notes"],
            })

        # ====================================================
        # CANCELLED ORDERS STOP HERE
        # ====================================================
        # Cancellation is already an operational processing route.
        # Exceptions are classified only after all processing sheets.
        # ====================================================

        if r["cancelled"]:
            continue

        # ====================================================
        # RETURN-REQUEST ORDERS
        # ====================================================
        # Return work does NOT block other applicable operational
        # work. A return can be in progress while fulfilled physical
        # items still require tracking and delivery.
        # ====================================================

        if r["return_requested"]:

            if (
                r["refund_completed"]
                and r["outstanding"] == 0
            ):

                sheets["01_No_Action_Required"].append({
                    "Order": r["order"],
                    "Reason": "Return requested and refund recorded",
                    "Financial Status": r["fin"],
                    "Fulfillment Status": r["ful"],
                    "Refunded Amount": r["refund_amount"],
                    "Outstanding Balance": r["outstanding"],
                    "Required Processing":
                        "Return/refund resolved; other applicable "
                        "fulfilment/logistics work may continue",
                    "Archive": "No",
                })

            # Deliberately no continue:
            # applicable fulfilment/tracking/delivery continues.

        # ====================================================
        # 07 FULFILMENT
        # ====================================================

        if (
            r["needs_fulfilment"]
            and not r["refund_pending"]
        ):

            for _, line in r["g"].iterrows():

                product = txt(
                    line.get(
                        "Lineitem name",
                        ""
                    )
                )

                if not product:
                    continue

                sheets["07_Fulfilment"].append({
                    "Order": r["order"],
                    "Current Fulfilment Status":
                        r["ful"],
                    "Product":
                        product,
                    "SKU":
                        txt(
                            line.get(
                                "Lineitem sku",
                                ""
                            )
                        ),
                    "Quantity":
                        num(
                            line.get(
                                "Lineitem quantity",
                                0
                            )
                        ),
                    "Requires Shipping":
                        txt(
                            line.get(
                                "Lineitem requires shipping",
                                ""
                            )
                        ),
                    "Location":
                        r["location"],
                    "Fulfilment Confirmed At":
                        display_dt(tracking_at),
                })

        # ====================================================
        # 08 TRACKING
        # ====================================================

        if (
            r["shipping_required"]
            and not r["refund_pending"]
        ):

            sheets["08_Tracking"].append({
                "Order": r["order"],
                "Courier": COURIER,
                "Tracking Number": tracking_number,
                "Tracking Assigned At":
                    display_dt(tracking_at),
            })

            # =================================================
            # 09 DELIVERY
            # =================================================

            sheets["09_Delivery"].append({
                "Order": r["order"],
                "Tracking Number": tracking_number,
                "Courier": COURIER,
                "Delivery Status": "Delivered",
                "Delivered At":
                    display_dt(delivery_at),
            })

    # ========================================================
    # EXCEPTIONS-LAST + SOURCE-AWARE IMPORT-ERROR RESOLUTION
    # ========================================================
    #
    # HARD RULE 1:
    # Any order already being processed in sheets 00-09 cannot
    # appear in the exception workflow.
    #
    # HARD RULE 2:
    # For a remaining exception candidate, inspect Source.
    #
    # If Source != "shopify_draft_order":
    # organization instruction treats it as an imported-data
    # error case (e.g. StoreRobo/Matrixify-originated or another
    # non-draft source). Route it to bulk corrective processing:
    #
    #   Mark Paid       = Yes
    #   Mark Fulfilled  = Yes
    #   Mark Delivered  = Yes
    #   Archive         = Yes
    #
    # Such an order is NOT retained in Exceptions because the
    # organization has supplied a direct resolution.
    #
    # If Source == "shopify_draft_order":
    # no import-error resolution is assumed; the unresolved order
    # remains in Exceptions for review.
    # ========================================================

    processing_sheet_names = [
        "00_Incorrect_Orders",
        "01_No_Action_Required",
        "02_Cancelled_Orders",
        "03_Return_Requested",
        "04_Payment_Received",
        "05_Refund_Processed",
        "06_Refund_Completed",
        "07_Fulfilment",
        "08_Tracking",
        "09_Delivery",
    ]

    processed_orders = set()

    for sheet_name in processing_sheet_names:
        for row in sheets[sheet_name]:
            order = txt(row.get("Order", ""))
            if order:
                processed_orders.add(order)

    # --------------------------------------------------------
    # Determine exception candidates only from orders that have
    # NO normal operational processing route.
    # --------------------------------------------------------

    exception_candidates = []

    for r in records:

        if r["order"] in processed_orders:
            continue

        if (
            not r["created"]
            or not r["currency"]
        ):
            exception_candidates.append((
                r,
                "E08",
                "Required source data missing",
                "Correct source data before operational processing",
            ))
            continue

        if (
            r["fin"] == "paid"
            and r["outstanding"] > 0
        ):
            exception_candidates.append((
                r,
                "E01",
                "Paid status with unexplained positive outstanding balance",
                "Verify historical financial state",
            ))
            continue

        if (
            r["fin"] in {"pending", "partially_paid"}
            and r["outstanding"] <= 0
            and not r["refund_pending"]
        ):
            exception_candidates.append((
                r,
                "E02",
                "Payment state inconsistent with balance",
                "Verify payment state",
            ))
            continue

        exception_candidates.append((
            r,
            "E10",
            "No direct processing route identified",
            "Review source order and determine required action",
        ))

    # --------------------------------------------------------
    # Organization-directed source resolution
    # --------------------------------------------------------

    for r, code, reason, action in exception_candidates:

        source_norm = norm(r["source"])

        if source_norm != "shopify_draft_order":

            sheets["10_Import_Error_Resolution"].append({
                "Order": r["order"],
                "Source": r["source"],
                "Original Financial Status": r["fin"],
                "Original Fulfillment Status": r["ful"],
                "Outstanding Balance": r["outstanding"],
                "Refunded Amount": r["refund_amount"],
                "Total": r["total"],
                "Original Exception Code": code,
                "Original Exception Type": reason,
                "Organization Classification":
                    "Imported-data error resolution",
                "Mark Paid": "Yes",
                "Mark Fulfilled": "Yes",
                "Mark Delivered": "Yes",
                "Archive": "Yes",
                "Required Action":
                    "Bulk mark paid, fulfilled, delivered and archived",
            })

        else:

            add_exception(
                sheets["11_Exceptions"],
                r,
                code,
                reason,
                action,
            )

    import_error_resolved_orders = {
        txt(row.get("Order", ""))
        for row in sheets["10_Import_Error_Resolution"]
        if txt(row.get("Order", ""))
    }

    # ========================================================
    # COVERAGE AUDIT
    # ========================================================

    membership = {
        r["order"]: []
        for r in records
    }

    for sheet_name, rows in sheets.items():

        for row in rows:

            order = row.get("Order", "")

            if (
                order in membership
                and sheet_name not in membership[order]
            ):
                membership[order].append(
                    sheet_name
                )

    # ========================================================
    # MUTUAL-EXCLUSIVITY SAFETY CHECK
    # ========================================================

    exception_orders = {
        txt(row.get("Order", ""))
        for row in sheets["11_Exceptions"]
        if txt(row.get("Order", ""))
    }

    overlap_normal_exception = exception_orders & processed_orders
    overlap_import_exception = exception_orders & import_error_resolved_orders
    overlap_normal_import = processed_orders & import_error_resolved_orders

    if overlap_normal_exception:
        raise RuntimeError(
            "Internal classification error: order(s) appear in both "
            "normal processing and Exceptions: "
            + ", ".join(sorted(overlap_normal_exception))
        )

    if overlap_import_exception:
        raise RuntimeError(
            "Internal classification error: order(s) appear in both "
            "Import Error Resolution and Exceptions: "
            + ", ".join(sorted(overlap_import_exception))
        )

    if overlap_normal_import:
        raise RuntimeError(
            "Internal classification error: order(s) appear in both "
            "normal processing and Import Error Resolution: "
            + ", ".join(sorted(overlap_normal_import))
        )

    coverage = []

    for r in records:

        allocated = membership[
            r["order"]
        ]

        coverage.append({
            "Order": r["order"],
            "Total": r["total"],
            "Currency": r["currency"],
            "Source": r["source"],
            "Incorrect Order":
                "Yes"
                if r["incorrect_order"]
                else "No",
            "In Cancelled File":
                "Yes"
                if r["in_cancel_file"]
                else "No",
            "Return Requested":
                "Yes"
                if r["return_requested"]
                else "No",
            "Financial Status":
                r["fin"],
            "Fulfillment Status":
                r["ful"],
            "Outstanding Balance":
                r["outstanding"],
            "Refunded Amount":
                r["refund_amount"],
            "Shipping Required":
                "Yes"
                if r["shipping_required"]
                else "No",
            "Allocated Sheets":
                ", ".join(allocated),
            "Operational Processing":
                "Yes"
                if r["order"] in processed_orders
                else "No",
            "Import Error Resolution":
                "Yes"
                if r["order"] in import_error_resolved_orders
                else "No",
            "Unresolved Exception":
                "Yes"
                if "11_Exceptions" in allocated
                else "No",
            "Covered":
                "Yes"
                if allocated
                else "No",
        })

    sheets["11_Order_Coverage"] = coverage

    # ========================================================
    # CREATE EXCEL
    # ========================================================

    wb = Workbook()

    wb.remove(
        wb.active
    )

    for sheet_name, rows in sheets.items():

        ws = wb.create_sheet(
            sheet_name
        )

        if rows:

            headers = list(
                rows[0].keys()
            )

            ws.append(
                headers
            )

            for row in rows:

                ws.append([
                    row.get(h, "")
                    for h in headers
                ])

        else:

            ws.append([
                "No matching orders"
            ])

        # HEADER
        for cell in ws[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="D9EAF7"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # WIDTH
        for col in range(
            1,
            ws.max_column + 1
        ):

            width = 12

            for row in range(
                1,
                min(
                    ws.max_row,
                    100
                ) + 1
            ):

                value = ws.cell(
                    row=row,
                    column=col
                ).value

                width = max(
                    width,
                    len(
                        str(
                            value or ""
                        )
                    ) + 2
                )

            ws.column_dimensions[
                get_column_letter(col)
            ].width = min(
                width,
                45
            )

        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    # ========================================================
    # SAVE
    # ========================================================

    now = datetime.now()

    output = Path(output_path) if output_path else (
        Path.cwd()
        / (
            "inkor_fmcg_worksheet_"
            f"{now:%Y%m%d_%H%M%S}.xlsx"
        )
    )

    wb.save(
        output
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    print()
    print("=" * 75)
    print("INKOR FMCG BACKLOG WORKSHEET GENERATED")
    print("=" * 75)

    print(
        f"Main orders CSV       : {main_csv}"
    )

    print(
        f"Return-request CSV    : {return_csv}"
    )

    print(
        f"Cancelled-orders CSV  : {cancelled_csv}"
    )

    print("-" * 75)

    print(
        f"Unique main orders    : {len(records)}"
    )

    print(
        f"Return-request orders : {len(return_orders)}"
    )

    print(
        f"Cancelled orders      : {len(cancelled_orders)}"
    )

    print("-" * 75)

    for sheet_name, rows in sheets.items():

        print(
            f"{sheet_name:27s}: "
            f"{len(rows)} rows"
        )

    print("-" * 75)

    print(
        f"Output: {output}"
    )

    print("=" * 75)

    return output




# ============================================================
# STORE ROBO 201-COLUMN FORMAT
# Exact column order taken from the supplied StoreRobo CSV.
# ============================================================

STOREROBO_HEADERS = [
    "ID","Name","Send Receipt","Inventory Behaviour","Number","Phone","Email",
    "Note","Tags","Created At","Updated At","Cancelled At","Cancel: Reason",
    "Cancel: Send Receipt","Cancel: Refund","Processed At","Closed At","Currency",
    "Source","Source Identifier","Source URL","User ID","Order Status URL",
    "Weight Total","Price: Total Line Items","Price: Current Subtotal",
    "Price: Subtotal","Tax 1: Title","Tax 1: Rate","Tax 1: Price","Tax 2: Title",
    "Tax 2: Rate","Tax 2: Price","Tax 3: Title","Tax 3: Rate","Tax 3: Price",
    "Tax: Included","Tax: Total","Price: Total Discount","Price: Total Shipping",
    "Price: Current Total Duties","Price: Total Duties","Price: Current Total Fees",
    "Price: Total Fees","Price: Total Refund","Price: Total Outstanding",
    "Price: Current Total","Price: Total","Payment: Status",
    "Order Fulfillment Status","Archived","Purchase Order Number","Customer: ID",
    "Customer: Email","Customer: Phone","Customer: First Name","Customer: Last Name",
    "Customer: Note","Customer: State","Customer: Tags",
    "Customer: Email Marketing Status","Customer: SMS Marketing Status",
    "Customer: Tax Exempt","Customer: Orders Count","Customer: Total Spent",
    "Billing: First Name","Billing: Last Name","Billing: Name","Billing: Company",
    "Billing: Phone","Billing: Address 1","Billing: Address 2","Billing: Zip",
    "Billing: City","Billing: Province","Billing: Province Code","Billing: Country",
    "Billing: Country Code","Shipping: First Name","Shipping: Last Name",
    "Shipping: Name","Shipping: Company","Shipping: Phone","Shipping: Address 1",
    "Shipping: Address 2","Shipping: Zip","Shipping: City","Shipping: Province",
    "Shipping: Province Code","Shipping: Country","Shipping: Country Code",
    "Company: ID","Company: Name","Company: External ID","Company: Location ID",
    "Company: Location Name","Company: Location External ID","Line: Type","Line: ID",
    "Line: Product ID","Line: Product Handle","Line: Title","Line: Name",
    "Line: Variant ID","Line: Variant Title","Line: SKU","Line: Quantity",
    "Line: Price","Line: Discount","Line: Discount Allocation","Line: Total",
    "Line: Requires Shipping","Line: Vendor","Line: Gift Card",
    "Line: Force Gift Card","Line: Taxable","Line: Tax Total","Line: Tax 1 Title",
    "Line: Tax 1 Rate","Line: Tax 1 Price","Line: Tax 2 Title","Line: Tax 2 Rate",
    "Line: Tax 2 Price","Line: Tax 3 Title","Line: Tax 3 Rate","Line: Tax 3 Price",
    "Line: Fulfillable Quantity","Line: Fulfillment Service",
    "Line: Fulfillment Status","Shipping: ID","Shipping: Title","Shipping: Code",
    "Shipping: Carrier Identifier","Shipping: Source","Shipping: Price",
    "Shipping: Tax Total","Shipping: Tax 1 Title","Shipping: Tax 1 Rate",
    "Shipping: Tax 1 Price","Shipping: Tax 2 Title","Shipping: Tax 2 Rate",
    "Shipping: Tax 2 Price","Discount: Type","Discount: Code","Discount: Value",
    "Discount: Percentage","Line: Product Type","Line: Product Tags",
    "Line: Variant SKU","Line: Variant Barcode","Line: Variant Weight",
    "Line: Variant Weight Unit","Line: Variant Inventory Qty","Line: Variant Cost",
    "Line: Variant Price","Line: Variant Compare At Price",
    "Line: Variant Country of Origin","Line: Variant Province of Origin",
    "Line: Variant HS Code","Refund: ID","Refund: Created At","Refund: Note",
    "Refund: Restock","Refund: Restock Type","Refund: Restock Location",
    "Refund: Send Receipt","Refund: Generate Transaction","Refund Adjustment: ID",
    "Refund Adjustment: Amount","Transaction: ID","Transaction: Kind",
    "Transaction: Processed At","Transaction: Amount","Transaction: Currency",
    "Transaction: Status","Transaction: Message","Transaction: Gateway",
    "Transaction: Force Gateway","Transaction: Test","Transaction: Authorization",
    "Transaction: Payment ID","Transaction: Device ID","Transaction: User ID",
    "Transaction: Parent ID","Transaction: Error Code","Transaction: CC AVS Result",
    "Transaction: CC Bin","Transaction: CC CVV Result","Transaction: CC Number",
    "Transaction: CC Company","Fulfillment: ID","Fulfillment: Status",
    "Fulfillment: Created At","Fulfillment: Updated At",
    "Fulfillment: Tracking Company","Fulfillment: Location",
    "Fulfillment: Shipment Status","Fulfillment: Processed At",
    "Fulfillment: Tracking Number","Fulfillment: Tracking URL",
    "Fulfillment: Send Receipt"
]


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


def normalize_order(v):
    return clean(v)


def read_order_export(path):
    """
    Read the standard Shopify order export.

    One Shopify order can occupy multiple CSV rows.  For this task we only
    need order-level information needed by the StoreRobo fulfillment row.

    User-defined rule:
        StoreRobo ID = Shopify export Name
    """
    orders = {}

    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)

        if not reader.fieldnames:
            raise ValueError("Order export CSV has no header row.")

        required = {"Name", "Location"}
        missing = required - set(reader.fieldnames)

        if missing:
            raise ValueError(
                "Order export CSV is missing required column(s): "
                + ", ".join(sorted(missing))
            )

        for row in reader:
            name = normalize_order(row.get("Name"))

            if not name:
                continue

            if name not in orders:
                orders[name] = {
                    "Name": name,
                    "Location": clean(row.get("Location")),
                }
            else:
                # Use the first nonblank location found for the order.
                if not orders[name]["Location"]:
                    orders[name]["Location"] = clean(row.get("Location"))

    return orders


def sheet_rows_as_dicts(ws):
    rows = ws.iter_rows(values_only=True)

    try:
        headers = [clean(v) for v in next(rows)]
    except StopIteration:
        return []

    result = []

    for values in rows:
        row = {
            headers[i]: values[i]
            for i in range(min(len(headers), len(values)))
            if headers[i]
        }

        if any(clean(v) for v in values):
            result.append(row)

    return result


def read_workbook(path):
    """
    Read tracking and delivery instructions from the generated worksheet.

    Required sheets:
        08_Tracking
        09_Delivery
    """
    wb = load_workbook(path, read_only=True, data_only=True)

    required_sheets = {"08_Tracking", "09_Delivery"}
    missing = required_sheets - set(wb.sheetnames)

    if missing:
        raise ValueError(
            "Worksheet is missing required sheet(s): "
            + ", ".join(sorted(missing))
        )

    tracking_rows = sheet_rows_as_dicts(wb["08_Tracking"])
    delivery_rows = sheet_rows_as_dicts(wb["09_Delivery"])

    tracking = {}
    delivery = {}

    for row in tracking_rows:
        order = normalize_order(row.get("Order"))

        if not order:
            continue

        tracking[order] = {
            "Courier": clean(row.get("Courier")),
            "Tracking Number": clean(row.get("Tracking Number")),
            "Tracking Assigned At": clean(row.get("Tracking Assigned At")),
        }

    for row in delivery_rows:
        order = normalize_order(row.get("Order"))

        if not order:
            continue

        delivery[order] = {
            "Tracking Number": clean(row.get("Tracking Number")),
            "Courier": clean(row.get("Courier")),
            "Delivery Status": clean(row.get("Delivery Status")),
            "Delivered At": clean(row.get("Delivered At")),
        }

    return tracking, delivery


def make_blank_storerobo_row():
    return {h: "" for h in STOREROBO_HEADERS}


def build_storerobo_row(order, order_info, tracking_info, delivery_info):
    """
    Create one StoreRobo Fulfillment Line.

    Only fields supported by the two supplied inputs are populated.
    Unsupported StoreRobo fields remain blank.

    IMPORTANT:
    Per the user's course-data rule:
        ID = Name
    """
    row = make_blank_storerobo_row()

    order_name = order
    courier = clean(tracking_info.get("Courier"))
    tracking_number = clean(tracking_info.get("Tracking Number"))
    tracking_at = clean(tracking_info.get("Tracking Assigned At"))

    delivered_status = clean(delivery_info.get("Delivery Status"))
    delivered_at = clean(delivery_info.get("Delivered At"))

    # Prefer tracking sheet values. Delivery sheet is a consistency fallback.
    if not courier:
        courier = clean(delivery_info.get("Courier"))

    if not tracking_number:
        tracking_number = clean(delivery_info.get("Tracking Number"))

    # --------------------------------------------------------
    # ORDER IDENTIFICATION
    # --------------------------------------------------------
    row["ID"] = order_name
    row["Name"] = order_name

    # --------------------------------------------------------
    # STORE ROBO ROW TYPE
    # --------------------------------------------------------
    row["Line: Type"] = "Fulfillment Line"

    # --------------------------------------------------------
    # FULFILLMENT / TRACKING
    # --------------------------------------------------------
    row["Fulfillment: Status"] = "SUCCESS"
    row["Fulfillment: Created At"] = tracking_at
    row["Fulfillment: Updated At"] = tracking_at
    row["Fulfillment: Tracking Company"] = courier
    row["Fulfillment: Location"] = clean(order_info.get("Location"))
    row["Fulfillment: Tracking Number"] = tracking_number
    row["Fulfillment: Send Receipt"] = "FALSE"

    # --------------------------------------------------------
    # DELIVERY
    # --------------------------------------------------------
    if delivered_status:
        row["Fulfillment: Shipment Status"] = delivered_status.lower()

    if delivered_at:
        row["Fulfillment: Processed At"] = delivered_at

    return row


def validate_tracking_delivery(tracking, delivery):
    """
    Tracking is the controlling list.

    Delivery rows may exist for the same orders.  We do not create extra
    StoreRobo rows from delivery-only entries because the requested upload
    is the tracking-number import.
    """
    problems = []

    for order, t in tracking.items():
        if not clean(t.get("Tracking Number")):
            problems.append(f"{order}: tracking number is blank")

        if not clean(t.get("Courier")):
            problems.append(f"{order}: courier is blank")

    if problems:
        raise ValueError(
            "Tracking worksheet contains invalid row(s):\n  - "
            + "\n  - ".join(problems)
        )


def generate_tracking_upload(order_csv, workbook_path, output_path=None):
    orders = read_order_export(order_csv)
    tracking, delivery = read_workbook(workbook_path)

    validate_tracking_delivery(tracking, delivery)

    output_rows = []
    missing_orders = []

    for order, tracking_info in tracking.items():

        if order not in orders:
            missing_orders.append(order)
            continue

        delivery_info = delivery.get(order, {})

        # If both sheets carry a tracking number, they must agree.
        tnum = clean(tracking_info.get("Tracking Number"))
        dnum = clean(delivery_info.get("Tracking Number"))

        if tnum and dnum and tnum != dnum:
            raise ValueError(
                f"{order}: tracking number differs between "
                f"08_Tracking ({tnum}) and 09_Delivery ({dnum})."
            )

        output_rows.append(
            build_storerobo_row(
                order,
                orders[order],
                tracking_info,
                delivery_info,
            )
        )

    if missing_orders:
        raise ValueError(
            "The following worksheet tracking order(s) are absent from "
            "the Shopify order export: "
            + ", ".join(missing_orders)
        )

    if not output_rows:
        raise ValueError("No tracking rows were found to export.")

    now = datetime.now()
    output = Path(output_path) if output_path else (
        Path.cwd()
        / f"tracking_upload_{now:%Y%m%d_%H%M%S}.csv"
    )

    with open(output, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=STOREROBO_HEADERS,
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(output_rows)

    print()
    print("=" * 72)
    print("STORE ROBO TRACKING UPLOAD CREATED")
    print("=" * 72)
    print(f"Order export       : {order_csv}")
    print(f"INKOR worksheet    : {workbook_path}")
    print(f"Tracking orders    : {len(tracking)}")
    print(f"Delivery records   : {len(delivery)}")
    print(f"Output rows        : {len(output_rows)}")
    print(f"StoreRobo columns  : {len(STOREROBO_HEADERS)}")
    print(f"Output             : {output}")
    print("=" * 72)
    print()

    return output




def run_app():
    st.set_page_config(
        page_title="INKOR Fulfilment Data Preparation Tool",
        page_icon="📦",
        layout="centered",
    )

    st.title("INKOR Fulfilment Data Preparation Tool")
    st.caption("Lecture 07 • Upload the three Shopify order exports to prepare your operational files.")

    with st.expander("Files required", expanded=True):
        st.markdown(
            "1. **Complete order export**\
"
            "2. **Return-request order export**\
"
            "3. **Cancelled-order export**"
        )

    main_file = st.file_uploader("1. Complete order export", type=["csv"], key="main")
    return_file = st.file_uploader("2. Return-request order export", type=["csv"], key="return")
    cancelled_file = st.file_uploader("3. Cancelled-order export", type=["csv"], key="cancelled")

    ready = all([main_file, return_file, cancelled_file])

    if st.button("Generate Lecture 07 Files", type="primary", use_container_width=True, disabled=not ready):
        try:
            with st.spinner("Preparing operational files..."):
                with tempfile.TemporaryDirectory() as td:
                    td = Path(td)
                    main_path = td / "complete_orders.csv"
                    return_path = td / "return_requests.csv"
                    cancelled_path = td / "cancelled_orders.csv"
                    workbook_path = td / "lecture07_student_worksheet.xlsx"
                    tracking_path = td / "lecture07_tracking_upload.csv"

                    main_path.write_bytes(main_file.getvalue())
                    return_path.write_bytes(return_file.getvalue())
                    cancelled_path.write_bytes(cancelled_file.getvalue())

                    build_workbook(
                        main_path, return_path, cancelled_path,
                        output_path=workbook_path
                    )
                    generate_tracking_upload(
                        main_path, workbook_path,
                        output_path=tracking_path
                    )

                    st.session_state["worksheet_bytes"] = workbook_path.read_bytes()
                    st.session_state["tracking_bytes"] = tracking_path.read_bytes()

                    wb_check = load_workbook(workbook_path, read_only=True, data_only=True)
                    summary = {}
                    for sheet in wb_check.sheetnames:
                        ws = wb_check[sheet]
                        # Empty-result sheets contain only the generated header row.
                        summary[sheet] = max(ws.max_row - 1, 0)
                    wb_check.close()
                    st.session_state["summary"] = summary

            st.success("Files generated successfully.")
        except Exception as exc:
            st.session_state.pop("worksheet_bytes", None)
            st.session_state.pop("tracking_bytes", None)
            st.session_state.pop("summary", None)
            st.error(f"Could not generate the files: {exc}")

    if "worksheet_bytes" in st.session_state:
        st.subheader("Download")
        st.download_button(
            "Download Operational Worksheet",
            data=st.session_state["worksheet_bytes"],
            file_name="lecture07_student_worksheet.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
        st.download_button(
            "Download StoreRobo Tracking CSV",
            data=st.session_state["tracking_bytes"],
            file_name="lecture07_tracking_upload.csv",
            mime="text/csv",
            use_container_width=True,
        )

        with st.expander("Generation summary"):
            for name, count in st.session_state.get("summary", {}).items():
                st.write(f"{name}: {count} data rows")


if __name__ == "__main__":
    run_app()
